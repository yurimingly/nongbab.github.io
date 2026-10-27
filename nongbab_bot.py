"""농밥 봇 — 식당별 식단표 메일을 찾아 첨부(PDF·엑셀)를 읽고, 앱이 보여줄 JSON을 만듭니다.

흐름:  config.CAFETERIAS 의 식당마다 Gmail(IMAP) 검색 → 새 메일의 첨부
      → Claude가 표를 JSON으로 변환 → 형식 검증·정리
      → docs/data/<식당>/weeks/<월요일>.json 저장 → 목록·오프라인 사본 갱신

사용법:
  python nongbab_bot.py                          메일함 확인 후 새 식단표 반영 (GitHub Actions가 이걸 실행)
  python nongbab_bot.py --file 식단.xlsx --cafe center   메일 없이 로컬 파일 하나로 테스트
  python nongbab_bot.py --reprocess              이미 처리한 메일도 다시 읽기 (읽는 방식을 바꿨을 때)
  python nongbab_bot.py --reprocess --cafe center   한 식당만 다시 읽기
  python nongbab_bot.py --dump                   원본 첨부만 originals/mail/<식당>/ 에 내려받기 (대조용)
"""
import argparse
import base64
import datetime as dt
import email
import imaplib
import json
import os
import re
import sys
import urllib.error
import urllib.request
from email.header import decode_header, make_header
from email.utils import parsedate_to_datetime

import config

WEEKDAYS = "월화수목금토일"
IN_ACTIONS = os.environ.get("GITHUB_ACTIONS") == "true"


def report(level, msg):
    """콘솔에 찍고, GitHub Actions에서는 실행 화면 상단 알림 + 요약에도 남깁니다 (로그를 안 열어도 원인이 보이게)."""
    print(msg)
    if IN_ACTIONS:
        print(f"::{level}::" + msg.replace("%", "%25").replace("\r", "").replace("\n", "%0A"))
        summary = os.environ.get("GITHUB_STEP_SUMMARY")
        if summary:
            with open(summary, "a", encoding="utf-8") as f:
                f.write(("❌ " if level == "error" else "ℹ️ ") + msg.replace("\n", "  \n") + "\n\n")


class Fatal(Exception):
    """설정·로그인처럼 더 진행할 수 없는 문제"""

PROMPT = """첨부한 자료는 회사 구내식당의 주간 식단표입니다 (PDF와 확대 이미지, 또는 엑셀 표를 옮긴 텍스트).
표 전체를 빠짐없이 아래 JSON 형식으로 옮겨 주세요. JSON만 출력하고 설명은 붙이지 마세요.

{
  "title": "표 제목 그대로",
  "days": [
    {
      "date": "YYYY-MM-DD",
      "holiday": false,
      "holiday_name": "",
      "meals": [
        {
          "meal": "아침 | 점심 | 저녁 중 하나",
          "label": "표 왼쪽 끝에 적힌 구분 이름 그대로 (예: 활기찬 아침, 건강한 점심)",
          "time": "운영 시간 (예: 7:30~8:40)",
          "note": "그 끼니 전체에 대한 짧은 안내가 있으면 (없으면 빈 문자열)",
          "corners": [
            {"name": "코너 이름", "main": "대표 메뉴 (원산지)", "items": ["반찬 한 줄씩"], "kcal": 788,
             "time": "코너별 시간이 따로 있으면 (없으면 빈 문자열)", "note": "코너에 붙은 짧은 안내 (없으면 빈 문자열)"}
          ]
        }
      ]
    }
  ],
  "shared": [
    {"meal": "끼니", "name": "여러 요일에 걸쳐 합쳐진 칸의 이름 (예: 셀프라면)", "time": "운영 시간", "lines": ["한 줄씩"]}
  ]
}

표 읽는 법:
- 한 끼니 안에 코너(메뉴 선택지)가 여러 개 있습니다. 표 왼쪽 두 번째 칸이 코너 이름입니다 (예: CORNER1, 컵밥, PLUS, CORNER1 인터, 셀비 PLUS, 후식음료, TO샐러드(예약)).
- 코너마다 색칠된 칸(노란색·회색 등)이 맨 위에 있으면 그것이 그 코너의 대표 메뉴(main)이고, 그 아래 흰 칸의 메뉴들이 items입니다.
  대표 메뉴는 국·찌개만이 아니라 볶음·덮밥·우동 등 무엇이든 될 수 있습니다. 대표 메뉴를 items에 다시 넣지 마세요.
- 색칠된 대표 칸 없이 한 칸만 있는 코너(예: 컵밥 "돈가스컵밥 + 국", PLUS "구수한 숭늉", 후식음료 "녹차")는 그 칸 내용을 main에 넣고 items는 [].
- 코너 이름 칸이 비어 있는데 메뉴가 있으면 name은 빈 문자열 ""로 두세요. 이름을 지어내지 마세요.
- 칼로리는 그 코너 바로 아래 Kcal 칸의 숫자, 없으면 null.
- 원산지 괄호는 해당 메뉴 뒤에 그대로 붙이세요. 예: "모둠커틀렛&밥 (치킨가스-계육:국산/멘치가스-돈육:국산)"
- 한 칸 안에서 줄이 바뀌었어도 괄호가 이어지는 한 메뉴는 한 줄로 합치세요.
- 표에 있는 날짜는 모두 days에 넣으세요. 빈 칸이나 회색으로 막힌 날은 holiday=true, meals=[], 사유가 있으면 holiday_name.
- 메뉴 이름은 한 글자씩 확인해서 정확히 옮기세요 (예: '칩'과 '침', '진'과 '전'을 구분). 읽을 수 없는 글자는 [?].
- 엑셀 텍스트로 받은 경우: 셀 주소(예: C7)와 값이 주어지고, 병합된 셀은 범위로 표시됩니다. 같은 열이 같은 날짜입니다.
  [굵게]·[배경:색] 표시는 원본 서식이며, 보통 대표 메뉴나 구분 제목을 뜻합니다. 메뉴판 맨 위 안내 문구(캠페인 등)는 메뉴가 아닙니다.
"""

XLSX_NOTE = "아래는 엑셀 식단표의 셀 내용입니다 (형식: 셀주소 [서식] 값). 병합 셀은 '병합 범위' 목록을 참고하세요."


# ─────────────────────────── Gmail ───────────────────────────
def _decode(value):
    return str(make_header(decode_header(value))) if value else ""


def processed_key(cafe_id, msg_id):
    return f"{cafe_id}|{msg_id}"


def load_state():
    """처리한 메일 목록. 식당이 둘(손보·본관)이던 시절의 식당 구분 없는 기록은 그 두 식당 기록으로 옮김."""
    state = _load(config.PROCESSED_FILE, {"ids": []})
    ids = []
    for i in state.get("ids", []):
        if "|" in i:
            ids.append(i)
        else:
            ids += [processed_key("sonbo", i), processed_key("center", i)]
    return {"ids": list(dict.fromkeys(ids))}


class Mailbox:
    """Gmail에 한 번 로그인해서 식당마다 검색합니다."""

    def __init__(self):
        user = os.environ.get("GMAIL_USER")
        pw = os.environ.get("GMAIL_APP_PASSWORD", "").replace(" ", "")
        if not user or not pw:
            raise Fatal("GMAIL_USER / GMAIL_APP_PASSWORD 환경변수가 없습니다. README의 '1. Gmail 앱 비밀번호'를 보세요.")
        self.conn = imaplib.IMAP4_SSL(config.IMAP_HOST)
        try:
            self.conn.login(user, pw)
        except imaplib.IMAP4.error as e:
            raise Fatal(f"Gmail 로그인 실패: {e}\n→ 일반 비밀번호가 아니라 '앱 비밀번호'를 넣었는지 확인하세요.")
        # 받은편지함 밖(라벨·보관함)에 있어도 찾도록 전체보관함에서 검색.
        # 한국어 Gmail은 폴더 이름이 '[Gmail]/전체보관함'이라, 이름 대신 \All 표시로 찾습니다.
        all_mail = None
        for line in self.conn.list()[1] or []:
            text = line.decode(errors="ignore") if isinstance(line, bytes) else str(line)
            if "\\All" in text:
                all_mail = text.rsplit(' "/" ', 1)[-1].strip()
                break
        if not all_mail or self.conn.select(all_mail, readonly=True)[0] != "OK":
            self.conn.select("INBOX", readonly=True)

    def close(self):
        try:
            self.conn.logout()
        except Exception:
            pass

    def fetch(self, query, processed_ids, reprocess=False, cafe_id=""):
        """검색 조건에 맞는 메일 중 이 식당에서 아직 처리하지 않은 것의 첨부들을 오래된 순서로 돌려줍니다."""
        self.conn.literal = query.encode("utf-8")  # 한글 검색어는 literal로 넘겨야 Gmail 검색이 받습니다
        status, data = self.conn.uid("SEARCH", "CHARSET", "UTF-8", "X-GM-RAW")
        if status != "OK":
            raise Fatal(f"메일 검색 실패 ({query}): {data}")
        all_uids = data[0].split()
        found = []
        for uid in all_uids[-config.MAX_MAILS_PER_RUN:]:
            status, msg_data = self.conn.uid("FETCH", uid, "(RFC822)")
            if status != "OK" or not msg_data or msg_data[0] is None:
                continue
            msg = email.message_from_bytes(msg_data[0][1])
            msg_id = (msg.get("Message-ID") or f"uid-{uid.decode()}").strip()
            if processed_key(cafe_id, msg_id) in processed_ids and not reprocess:
                continue
            try:
                received = parsedate_to_datetime(msg.get("Date")).isoformat()
            except Exception:
                received = ""
            files = []
            for part in msg.walk():
                filename = _decode(part.get_filename())
                ext = os.path.splitext(filename.lower())[1]
                if part.get_content_type() == "application/pdf" and not ext:
                    ext, filename = ".pdf", "menu.pdf"
                payload = part.get_payload(decode=True)
                if ext in config.ATTACHMENT_TYPES and payload:
                    files.append((filename, payload))
            found.append({"id": msg_id, "subject": _decode(msg.get("Subject")), "received": received, "files": files})
        found.sort(key=lambda m: m["received"])  # 오래된 것부터 → 같은 주 정정본이면 최신이 덮어씀
        return found, len(all_uids)


# ─────────────────────────── PDF → JSON ───────────────────────────
def page_images(pdf_bytes):
    """작은 글씨를 정확히 읽도록 PDF 첫 페이지를 고해상도로 그려 위·아래 두 장으로 나눕니다.
    pdftoppm(poppler)이 없으면 빈 목록 → PDF만으로 읽습니다."""
    import shutil, subprocess, tempfile
    if not shutil.which("pdftoppm"):
        return []
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "m.pdf")
        open(src, "wb").write(pdf_bytes)
        subprocess.run(["pdftoppm", "-r", str(config.RENDER_DPI), "-png", "-f", "1", "-l", "1", src,
                        os.path.join(tmp, "p")], check=True, capture_output=True)
        png = next(f for f in sorted(os.listdir(tmp)) if f.endswith(".png"))
        try:
            from PIL import Image
        except ImportError:
            return [open(os.path.join(tmp, png), "rb").read()]
        im = Image.open(os.path.join(tmp, png)).convert("RGB")
        w, h = im.size
        overlap = h // 30  # 경계에 걸친 줄이 잘리지 않게 겹쳐 자름
        out = []
        for box in [(0, 0, w, h // 2 + overlap), (0, h // 2 - overlap, w, h)]:
            path = os.path.join(tmp, f"part{len(out)}.png")
            im.crop(box).save(path, optimize=True)
            out.append(open(path, "rb").read())
        return out


def xlsx_text(data, sheets=None):
    """엑셀 식단표를 셀 주소·서식·값 목록으로 바꿉니다 (숨긴 행·열과 빈 셀 제외, 시트마다).
    sheets 가 주어지면 그 이름의 시트만 읽습니다."""
    import io
    try:
        import openpyxl
    except ImportError:
        raise Fatal("엑셀을 읽는 openpyxl 이 설치되지 않았습니다 (pip install openpyxl).")
    from openpyxl.utils import get_column_letter
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" or (sheets and ws.title not in sheets):
            continue
        hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
        hidden_cols = {c for c, d in ws.column_dimensions.items() if d.hidden}
        lines = []
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None or str(v).strip() == "" or cell.row in hidden_rows:
                    continue
                if get_column_letter(cell.column) in hidden_cols:
                    continue
                if isinstance(v, (dt.datetime, dt.date)):
                    v = v.strftime("%Y-%m-%d")
                tags = []
                if cell.font is not None and cell.font.bold:
                    tags.append("굵게")
                fill = cell.fill.fgColor.rgb if cell.fill is not None and cell.fill.fill_type == "solid" else None
                if isinstance(fill, str) and fill not in ("00000000", "FFFFFFFF"):
                    tags.append(f"배경:{fill[-6:]}")
                tag = f" [{','.join(tags)}]" if tags else ""
                text = str(v).replace("\r", "").replace("\n", " / ")
                lines.append(f"{cell.coordinate}{tag} {text}")
        if lines:
            merged = ", ".join(str(r) for r in ws.merged_cells.ranges)
            out.append(f"### 시트: {ws.title}\n병합 범위: {merged or '없음'}\n" + "\n".join(lines))
    if not out:
        raise ValueError(f"엑셀에 읽을 시트가 없습니다 (찾은 시트: {[w.title for w in wb.worksheets]}, 설정: {sheets})")
    return "\n\n".join(out)


def ask_claude(filename, data, hint="", sheets=None):
    key = os.environ.get("ANTHROPIC_API_KEY")
    if not key:
        raise Fatal("ANTHROPIC_API_KEY 환경변수가 없습니다. 저장소 Settings → Secrets에 넣어 주세요.")
    if filename.lower().endswith(".xlsx"):
        content = [{"type": "text", "text": XLSX_NOTE + "\n\n" + xlsx_text(data, sheets)}]
    else:
        content = [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                   "data": base64.b64encode(data).decode()}}]
        for i, img in enumerate(page_images(data)):
            content.append({"type": "text", "text": f"확대 이미지 {i + 1} ({'위쪽' if i == 0 else '아래쪽'} 절반)"})
            content.append({"type": "image", "source": {"type": "base64", "media_type": "image/png",
                                                        "data": base64.b64encode(img).decode()}})
    content.append({"type": "text", "text": PROMPT + (f"\n이 식당 표의 특징:\n{hint}\n" if hint else "")})
    body = {"model": config.CLAUDE_MODEL, "max_tokens": config.CLAUDE_MAX_TOKENS,
            "messages": [{"role": "user", "content": content}]}
    req = urllib.request.Request(
        "https://api.anthropic.com/v1/messages",
        data=json.dumps(body).encode(),
        headers={"x-api-key": key, "anthropic-version": "2023-06-01", "content-type": "application/json"},
    )
    try:
        with urllib.request.urlopen(req, timeout=180) as r:
            res = json.load(r)
    except urllib.error.HTTPError as e:
        detail = e.read().decode(errors="ignore")[:300]
        if e.code in (401, 403):
            raise Fatal(f"Claude API 키가 거절됐습니다 ({e.code}). 키를 다시 확인하세요: {detail}")
        if e.code == 400 and "credit" in detail.lower():
            raise Fatal(f"Claude API 잔액이 부족합니다. console.anthropic.com에서 충전하세요: {detail}")
        raise RuntimeError(f"Claude API 오류 {e.code}: {detail}")
    return "".join(b.get("text", "") for b in res.get("content", []) if b.get("type") == "text")


def extract_json(text):
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.M).strip()
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end < 0:
        raise ValueError("응답에 JSON이 없습니다")
    return json.loads(text[start:end + 1])


def _date(s):
    return dt.date.fromisoformat(str(s).strip()[:10])


def _kcal(v):
    if v in (None, ""):
        return None
    m = re.search(r"\d+", str(v).replace(",", ""))
    return int(m.group()) if m else None


def _lines(v):
    if isinstance(v, str):
        v = v.split("\n")
    out = []
    for x in v or []:
        t = re.sub(r"\s+", " ", str(x)).strip()
        t = re.sub(r"^[*•·\-–]+\s*", "", t)  # 표의 글머리 기호(*, •, -)는 앱이 따로 그리므로 제거
        if t:
            out.append(t)
    return out


MEAL_ORDER = {"아침": 0, "점심": 1, "저녁": 2}


def _meal_name(m, label=""):
    for name in MEAL_ORDER:
        if name in (m or "") or name in (label or ""):
            return name
    return (m or label or "식사").strip()


def normalize(raw, source):
    """Claude가 돌려준 값을 앱이 믿고 쓸 수 있는 형태로 정리·검증합니다. 문제가 있으면 ValueError."""
    days = []
    for d in raw.get("days") or []:
        date = _date(d["date"])
        meals = []
        for m in d.get("meals") or []:
            corners = []
            for c in m.get("corners") or []:
                items = _lines(c.get("items"))
                main = " ".join(_lines([c.get("main") or ""]))
                if not main and items:          # 대표 칸이 따로 없던 코너
                    main, items = items[0], items[1:]
                if items and items[0] == main:  # 대표 메뉴를 반찬에 또 넣은 경우
                    items = items[1:]
                if main:
                    corners.append({"name": str(c.get("name") or "").strip(), "main": main,
                                    "items": items, "kcal": _kcal(c.get("kcal")),
                                    "time": str(c.get("time") or "").strip(),
                                    "note": str(c.get("note") or "").strip()})
            note = str(m.get("note") or "").strip()
            if corners or note:
                meals.append({"meal": _meal_name(m.get("meal"), m.get("label")),
                              "label": str(m.get("label") or "").strip(),
                              "time": str(m.get("time") or "").strip(), "note": note, "corners": corners})
        meals.sort(key=lambda m: MEAL_ORDER.get(m["meal"], 9))
        days.append({"date": date.isoformat(), "weekday": WEEKDAYS[date.weekday()],
                     "holiday": bool(d.get("holiday")) or not any(m["corners"] for m in meals),
                     "holiday_name": str(d.get("holiday_name") or "").strip(), "meals": meals})
    if not days:
        raise ValueError("요일 데이터가 비어 있습니다")
    days.sort(key=lambda d: d["date"])
    if all(d["holiday"] for d in days):
        raise ValueError("모든 날이 비어 있습니다 — 표를 읽지 못한 것 같습니다")

    first = _date(days[0]["date"])
    week_start = (first - dt.timedelta(days=first.weekday())).isoformat()  # 항상 그 주 월요일 기준
    shared = [{"meal": _meal_name(s.get("meal")), "name": str(s.get("name") or "").strip(),
               "time": str(s.get("time") or "").strip(), "lines": _lines(s.get("lines"))}
              for s in raw.get("shared") or [] if _lines(s.get("lines"))]
    return {"title": str(raw.get("title") or "주간식단표").strip(),
            "week_start": week_start, "week_end": days[-1]["date"],
            "days": days, "shared": shared, "source": source,
            "updated": dt.datetime.now(dt.timezone(dt.timedelta(hours=9))).isoformat(timespec="seconds")}


def parse_file(cafe, filename, data, source):
    last_err = None
    for attempt in range(1 + config.PARSE_RETRIES):
        try:
            return normalize(extract_json(ask_claude(filename, data, cafe.get("hint", ""), cafe.get("sheets"))), source)
        except (ValueError, KeyError, json.JSONDecodeError) as e:
            last_err = e
            print(f"  ↻ 해석 재시도 {attempt + 1}: {e}")
    raise RuntimeError(f"식단표 해석 실패: {last_err}")


# ─────────────────────────── 저장 ───────────────────────────
def _load(path, default):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return default


def _save(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(path)  # 앱이 반쯤 쓰인 파일을 읽지 않도록 한 번에 교체


def cafe_dir(cafe):
    return config.DATA_DIR / cafe["id"]


def save_week(cafe, week, filename=None, data=None):
    _save(cafe_dir(cafe) / "weeks" / f"{week['week_start']}.json", week)
    if data:  # 읽은 결과와 언제든 대조할 수 있게 원본도 같은 이름으로 보관
        odir = config.ORIGINALS_DIR / cafe["id"]
        odir.mkdir(parents=True, exist_ok=True)
        ext = os.path.splitext(filename or "")[1].lower() or ".pdf"
        for old in odir.glob(f"{week['week_start']}.*"):
            old.unlink()
        (odir / f"{week['week_start']}{ext}").write_bytes(data)
    print(f"  ✓ [{cafe['name']}] {week['week_start']} ~ {week['week_end']} 저장")


def migrate_layout():
    """식당이 하나뿐이던 시절 구조(data/weeks, originals/*.pdf)를 손보밥 폴더로 옮깁니다. 여러 번 실행해도 안전."""
    first = config.CAFETERIAS[0]
    old_weeks = config.DATA_DIR / "weeks"
    if old_weeks.is_dir():
        new = cafe_dir(first) / "weeks"
        new.mkdir(parents=True, exist_ok=True)
        for f in old_weeks.glob("*.json"):
            f.replace(new / f.name)
        old_weeks.rmdir()
    (config.DATA_DIR / "index.json").unlink(missing_ok=True)
    odir = config.ORIGINALS_DIR / first["id"]
    for f in list(config.ORIGINALS_DIR.glob("*.pdf")):
        odir.mkdir(parents=True, exist_ok=True)
        f.replace(odir / f.name)


def rebuild_index():
    """식당별 weeks/ 폴더를 기준으로 목록·식당 목록·오프라인 사본을 다시 만듭니다 (파일과 목록이 어긋나지 않게)."""
    now = dt.datetime.now(dt.timezone(dt.timedelta(hours=9))).isoformat(timespec="seconds")
    bundle, cafes = {}, []
    for cafe in config.CAFETERIAS:
        d = cafe_dir(cafe)
        files = sorted((d / "weeks").glob("*.json"), reverse=True)
        for old in files[config.KEEP_WEEKS:]:
            old.unlink()
        weeks = []
        for f in files[:config.KEEP_WEEKS]:
            w = _load(f, None)
            if w:
                weeks.append({"start": w["week_start"], "end": w["week_end"], "file": f"weeks/{f.name}"})
                bundle[f"{cafe['id']}/weeks/{f.name}"] = w
        index = {"id": cafe["id"], "name": cafe["name"], "place": cafe["place"],
                 "latest": weeks[0]["start"] if weeks else None, "weeks": weeks, "updated": now}
        _save(d / "index.json", index)
        bundle[f"{cafe['id']}/index.json"] = index
        cafes.append({"id": cafe["id"], "name": cafe["name"], "place": cafe["place"],
                      "latest": index["latest"], "weeks": len(weeks)})
    listing = {"app": config.APP_NAME, "cafeterias": cafes, "updated": now}
    _save(config.CAFETERIAS_FILE, listing)
    bundle["cafeterias.json"] = listing

    # 같은 데이터를 스크립트 한 개로도 저장 → 인터넷이 없거나 앱 안(file://)에서 열어도 식단이 보이게
    tmp = config.BUNDLE_FILE.with_suffix(".tmp")
    tmp.write_text("window.NONGBAB_BUNDLE=" + json.dumps(bundle, ensure_ascii=False, separators=(",", ":")) + ";\n",
                   encoding="utf-8")
    tmp.replace(config.BUNDLE_FILE)


def write_status(result, message):
    """앱의 '마지막 확인' 표시와 새로고침 버튼이 실행 완료를 알아채는 데 쓰는 파일 (매 실행마다 갱신)."""
    _save(config.STATUS_FILE, {"checked_at": dt.datetime.now(dt.timezone(dt.timedelta(hours=9))).isoformat(timespec="seconds"),
                               "result": result, "message": message})


# ─────────────────────────── 실행 ───────────────────────────
def pick_cafes(cafe_id):
    if not cafe_id:
        return config.CAFETERIAS
    cafes = [c for c in config.CAFETERIAS if c["id"] == cafe_id]
    if not cafes:
        raise Fatal(f"식당 id '{cafe_id}' 가 config.CAFETERIAS 에 없습니다: {[c['id'] for c in config.CAFETERIAS]}")
    return cafes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cafe", help="한 식당만 처리 (config.CAFETERIAS 의 id)")
    ap.add_argument("--file", help="메일 대신 로컬 식단표 파일(PDF·xlsx)로 테스트 (--cafe 필요)")
    ap.add_argument("--reprocess", action="store_true", help="이미 처리한 메일도 다시 해석")
    ap.add_argument("--dump", nargs="?", const=str(config.ORIGINALS_DIR / "mail"), metavar="DIR",
                    help="해석하지 않고 검색된 메일의 원본 첨부만 저장 (기본: originals/mail/<식당>/)")
    args = ap.parse_args()
    cafes = pick_cafes(args.cafe)
    migrate_layout()

    if args.file:
        if not args.cafe:
            raise Fatal("--file 은 --cafe 와 함께 써 주세요.")
        data = open(args.file, "rb").read()
        name = os.path.basename(args.file)
        save_week(cafes[0], parse_file(cafes[0], name, data, {"filename": name}), name, data)
        rebuild_index()
        return

    box = Mailbox()
    try:
        if args.dump:
            total = []
            for cafe in cafes:
                out = os.path.join(os.path.abspath(args.dump), cafe["id"])
                os.makedirs(out, exist_ok=True)
                mails, matched = box.fetch(cafe["query"], set(), reprocess=True)
                n = 0
                for m in mails:
                    for filename, data in m["files"]:
                        n += 1
                        safe = re.sub(r"[^0-9A-Za-z가-힣._-]", "_", filename)
                        with open(os.path.join(out, f"{m['received'][:10]}_{n:02d}_{safe}"), "wb") as f:
                            f.write(data)
                total.append(f"{cafe['name']} {n}개(메일 {matched}통)")
            report("notice", "원본 첨부 저장: " + ", ".join(total))
            return

        state = load_state()
        failed, saved, summary = 0, {}, []
        for cafe in cafes:
            mails, matched = box.fetch(cafe["query"], set(state["ids"]), args.reprocess, cafe["id"])
            if not mails:
                summary.append(f"{cafe['name']}: 새 메일 없음 (검색된 {matched}통 모두 반영됨)")
                continue
            for m in mails:
                print(f"● [{cafe['name']}] {m['received'][:16]}  {m['subject']}")
                if not m["files"]:
                    print("  - 읽을 수 있는 첨부(PDF·xlsx) 없음, 건너뜀")
                for filename, data in m["files"]:
                    try:
                        week = parse_file(cafe, filename, data,
                                          {"subject": m["subject"], "received": m["received"], "filename": filename})
                        save_week(cafe, week, filename, data)
                        saved.setdefault(cafe["name"], []).append(week["week_start"])
                    except Exception as e:  # 한 파일 실패가 다른 주차·식당 반영을 막지 않게
                        if isinstance(e, Fatal):
                            raise
                        failed += 1
                        report("error", f"[{cafe['name']}] {m['subject']} / {filename} 해석 실패: {e}")
                        break
                else:
                    key = processed_key(cafe["id"], m["id"])
                    if key not in state["ids"]:
                        state["ids"].append(key)  # 성공한 메일만 '처리됨' → 실패한 건 다음 실행에 재시도
            summary.append(f"{cafe['name']}: 메일 {len(mails)}통 처리")
        state["ids"] = state["ids"][-500:]
        _save(config.PROCESSED_FILE, state)
    finally:
        box.close()

    rebuild_index()
    if saved:
        parts = []
        for name, starts in saved.items():
            latest = max(starts)[5:].replace("-", "/")
            parts.append(f"{name} {latest}주" + (f" 외 {len(starts) - 1}주" if len(starts) > 1 else ""))
        msg = ", ".join(parts) + " 반영"
    else:
        msg = "새 식단표 없음"
    if failed:
        msg += f" · {failed}건 읽기 실패"
    write_status("error" if failed and not saved else "new" if saved else "none", msg)
    report("notice", " / ".join(summary))
    if failed:
        sys.exit(1)  # Actions에서 빨간 표시로 실패를 알 수 있게


if __name__ == "__main__":
    try:
        main()
    except Fatal as e:
        report("error", str(e))
        write_status("error", "확인 실패 — 설정을 확인해 주세요")
        sys.exit(1)
    except Exception as e:  # 예상 못 한 오류도 원인이 화면에 남게
        report("error", f"예상 못 한 오류: {type(e).__name__}: {e}")
        write_status("error", "확인 중 오류가 났어요")
        raise
